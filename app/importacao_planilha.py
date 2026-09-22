"""Leitura de planilhas (CSV/XLSX) enviadas pelo usuário para importação em lote.

Compartilhado entre os importadores de carteira (app.api.carteira) e de leads
(app.api.leads) -- extraído de app.api.carteira para não duplicar a lógica de
parsing (achado da Fase 4 do roadmap pós-auditoria de Leads, 03/09/2026).
"""

import csv
import io
import unicodedata

from fastapi import HTTPException
from openpyxl import load_workbook

TAMANHO_MAXIMO_IMPORTACAO = 5_000_000
# Achado baixo da auditoria da carteira de processos (Fase 10, 22/09/2026):
# só havia limite de tamanho de arquivo, não de linhas -- um CSV compacto de
# 5MB pode ter centenas de milhares de linhas, cada uma virando consulta(s)
# e insert num único request síncrono, sem paginação. Mesma ordem de
# grandeza do limite explícito de vincular-lote/atribuir-lote (max_length).
LINHAS_MAXIMAS_IMPORTACAO = 5000


def chave_coluna(texto: str) -> str:
    sem_acentos = "".join(c for c in unicodedata.normalize("NFKD", texto or "") if not unicodedata.combining(c))
    return "".join(ch for ch in sem_acentos.lower() if ch.isalnum())


def ler_planilha(conteudo: bytes, filename: str) -> list[dict[str, str]]:
    """Lê CSV ou XLSX e devolve uma lista de registros com chaves normalizadas."""
    nome = (filename or "").lower()
    linhas: list[list[str]] = []
    if nome.endswith(".xls"):
        raise HTTPException(400, "Formato .xls (Excel antigo) não é suportado. Salve como .xlsx ou .csv.")
    if nome.endswith((".xlsx", ".xlsm")):
        try:
            wb = load_workbook(io.BytesIO(conteudo), read_only=True, data_only=True)
        except Exception as exc:  # noqa: BLE001 - arquivo inválido enviado pelo usuário
            raise HTTPException(400, "Planilha Excel inválida ou corrompida.") from exc
        planilha = wb.active
        for linha in planilha.iter_rows(values_only=True):
            linhas.append(["" if celula is None else str(celula).strip() for celula in linha])
        wb.close()
    else:
        texto = None
        for codificacao in ("utf-8-sig", "latin-1"):
            try:
                texto = conteudo.decode(codificacao)
                break
            except UnicodeDecodeError:
                continue
        if texto is None:
            raise HTTPException(400, "Não foi possível ler o arquivo (codificação não suportada).")
        primeira_linha = texto.splitlines()[0] if texto.splitlines() else ""
        contagens = {sep: primeira_linha.count(sep) for sep in (";", "\t", ",")}
        delimitador = max(contagens, key=lambda sep: contagens[sep]) if any(contagens.values()) else ","
        for linha in csv.reader(io.StringIO(texto), delimiter=delimitador):
            linhas.append([campo.strip() for campo in linha])

    linhas = [linha for linha in linhas if any(linha)]
    if len(linhas) < 2:
        return []
    if len(linhas) - 1 > LINHAS_MAXIMAS_IMPORTACAO:
        raise HTTPException(
            413, f"Planilha com {len(linhas) - 1} linhas -- o máximo por importação é {LINHAS_MAXIMAS_IMPORTACAO}."
        )
    cabecalho = [chave_coluna(coluna) for coluna in linhas[0]]
    return [{cabecalho[i]: (linha[i] if i < len(linha) else "") for i in range(len(cabecalho))} for linha in linhas[1:]]


def valor_coluna(registro: dict[str, str], fragmentos: tuple[str, ...]) -> str | None:
    for chave, valor in registro.items():
        if valor.strip() and any(fragmento in chave for fragmento in fragmentos):
            return valor.strip()
    return None
