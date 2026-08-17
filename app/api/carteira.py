import csv
import io
import re
import unicodedata
from datetime import UTC, datetime
from typing import Annotated, Literal

from fastapi import APIRouter, Depends, File, Form, HTTPException, Query, Request, UploadFile
from openpyxl import load_workbook
from pydantic import BaseModel, Field, field_validator
from sqlalchemy import and_, case, func, or_, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.auth import UsuarioAutenticado, exigir_permissao, hash_ip
from app.cli.consolidar_situacoes_marcas import consolidar_situacao
from app.crm import obter_ou_criar_empresa
from app.database import get_session
from app.models import (
    EmpresaCRM,
    EventoAuditoria,
    HistoricoEtapaCarteira,
    Movimentacao,
    Processo,
    ProcessoMonitorado,
    RpiImportacao,
    TipoProcesso,
    Titular,
    UsuarioOperacoes,
    processo_titulares,
)
from app.normalization import normalizar_numero_processo
from app.proxy import cliente_ip

router = APIRouter(prefix="/v1/admin/carteira", tags=["processos monitorados"])
SessionDep = Annotated[AsyncSession, Depends(get_session)]
ViewDep = Annotated[UsuarioAutenticado, Depends(exigir_permissao("portfolio.view"))]
ManageDep = Annotated[UsuarioAutenticado, Depends(exigir_permissao("portfolio.manage"))]

ETAPAS_KANBAN: tuple[tuple[str, str], ...] = (
    ("triagem", "Novo / Triagem"),
    ("aguardando_documentos", "Aguardando documentos"),
    ("documentacao_gru", "Documentação e GRU"),
    ("protocolado", "Protocolado"),
    ("aguardando_inpi", "Aguardando INPI"),
    ("exigencia_recurso", "Exigência / Recurso"),
    ("deferido_concessao", "Deferido / Concessão"),
    ("encerrado", "Encerrado"),
)
GRUPOS_SITUACAO_INPI: tuple[tuple[str, str], ...] = (
    ("em_tramitacao", "Em tramitação"),
    ("exigencia", "Exigência"),
    ("sobrestado", "Sobrestado"),
    ("recurso", "Recurso / 2ª instância"),
    ("deferido", "Deferido"),
    ("registrado", "Registro concedido"),
    ("indeferido", "Indeferido"),
    ("encerrado", "Arquivado / Extinto"),
    ("revisar", "Revisar classificação"),
)
EtapaKanban = Literal[
    "triagem",
    "aguardando_documentos",
    "documentacao_gru",
    "protocolado",
    "aguardando_inpi",
    "exigencia_recurso",
    "deferido_concessao",
    "encerrado",
]


def _normalizar_busca(valor: str) -> str:
    sem_acentos = "".join(
        caractere
        for caractere in unicodedata.normalize("NFKD", valor.strip())
        if not unicodedata.combining(caractere)
    )
    return " ".join(sem_acentos.casefold().split())


def _expressao_procurador():
    return func.regexp_replace(
        func.trim(func.immutable_unaccent(func.lower(Processo.procurador))),
        r"\s+",
        " ",
        "g",
    )


def _expressao_grupo_situacao_inpi():
    codigo = Processo.situacao_normalizada
    return case(
        (codigo.in_(("publicada", "em_exame", "oposicao")), "em_tramitacao"),
        (codigo == "exigencia", "exigencia"),
        (codigo == "suspensa", "sobrestado"),
        (codigo.in_(("recurso", "recurso_decidido")), "recurso"),
        (codigo.in_(("deferida", "deferida_parcial")), "deferido"),
        (codigo == "registrada", "registrado"),
        (codigo == "indeferida", "indeferido"),
        (
            codigo.in_(("arquivada", "inexistente", "extinta", "cancelada")),
            "encerrado",
        ),
        else_="revisar",
    )


def _nome_grupo_situacao_inpi(chave: str) -> str:
    return dict(GRUPOS_SITUACAO_INPI).get(chave, "Revisar classificação")


def _validar_grupo_situacao_inpi(valor: str | None) -> str | None:
    valor = (valor or "").strip()
    if not valor:
        return None
    if valor not in dict(GRUPOS_SITUACAO_INPI):
        raise HTTPException(status_code=422, detail="Situação do INPI inválida.")
    return valor


def _grupo_situacao_valor(codigo: str | None) -> str:
    if codigo in {"publicada", "em_exame", "oposicao"}:
        return "em_tramitacao"
    if codigo == "exigencia":
        return "exigencia"
    if codigo == "suspensa":
        return "sobrestado"
    if codigo in {"recurso", "recurso_decidido"}:
        return "recurso"
    if codigo in {"deferida", "deferida_parcial"}:
        return "deferido"
    if codigo == "registrada":
        return "registrado"
    if codigo == "indeferida":
        return "indeferido"
    if codigo in {"arquivada", "inexistente", "extinta", "cancelada"}:
        return "encerrado"
    return "revisar"


def _titulo_exibicao(processo: Processo) -> str:
    titulo = (processo.titulo or "").strip()
    if titulo:
        return titulo
    if _normalizar_busca(processo.apresentacao or "") == "figurativa":
        return "Marca figurativa (sem elemento nominativo)"
    if processo.situacao_normalizada == "inexistente":
        return "Pedido inexistente — título não publicado"
    return "Título não informado pelo INPI"


class VinculoBase(BaseModel):
    empresa_id: int | None = Field(default=None, ge=1)
    empresa_nome: str | None = Field(default=None, min_length=2, max_length=200)
    responsavel_id: int | None = Field(default=None, ge=1)
    observacoes: str | None = Field(default=None, max_length=4000)

    @field_validator("empresa_nome", "observacoes", mode="before")
    @classmethod
    def limpar_opcionais(cls, valor: str | None) -> str | None:
        return valor.strip() or None if isinstance(valor, str) else valor


class CadastroManual(VinculoBase):
    numero: str = Field(min_length=5, max_length=50)


class VinculoLote(VinculoBase):
    processo_ids: list[int] = Field(min_length=1, max_length=200)


class VinculoProcurador(VinculoBase):
    procurador: str = Field(min_length=2, max_length=300)
    modo: Literal["exato", "contem", "variacoes"] = "exato"
    maximo: int = Field(default=5000, ge=1, le=5000)


class AtualizacaoMonitoramento(BaseModel):
    status: Literal["ativo", "pausado", "encerrado", "arquivado"] | None = None
    empresa_id: int | None = Field(default=None, ge=1)
    empresa_nome: str | None = Field(default=None, min_length=2, max_length=200)
    responsavel_id: int | None = Field(default=None, ge=1)
    remover_responsavel: bool = False
    observacoes: str | None = Field(default=None, max_length=4000)
    procurador: str | None = Field(default=None, max_length=500)
    etapa_kanban: EtapaKanban | None = None

    @field_validator("procurador", mode="before")
    @classmethod
    def _limpar_procurador(cls, valor: object) -> str | None:
        if valor is None:
            return None
        limpo = re.sub(r"\s+", " ", str(valor)).strip()
        return limpo or None


async def _empresa(
    session: AsyncSession,
    usuario: UsuarioAutenticado,
    empresa_id: int | None,
    empresa_nome: str | None,
) -> EmpresaCRM | None:
    if empresa_id is not None:
        empresa = (
            await session.execute(
                select(EmpresaCRM).where(
                    EmpresaCRM.id == empresa_id,
                    EmpresaCRM.organizacao_id == usuario.organizacao_id,
                )
            )
        ).scalar_one_or_none()
        if empresa is None:
            raise HTTPException(404, "Empresa não encontrada")
        return empresa
    return await obter_ou_criar_empresa(session, usuario.organizacao_id, empresa_nome)


async def _validar_responsavel(
    session: AsyncSession, usuario: UsuarioAutenticado, responsavel_id: int | None
) -> None:
    if responsavel_id is None:
        return
    existe = await session.scalar(
        select(UsuarioOperacoes.id).where(
            UsuarioOperacoes.id == responsavel_id,
            UsuarioOperacoes.organizacao_id == usuario.organizacao_id,
            UsuarioOperacoes.ativo.is_(True),
        )
    )
    if existe is None:
        raise HTTPException(404, "Responsável não encontrado")


def _auditar(
    session: AsyncSession,
    request: Request,
    usuario: UsuarioAutenticado,
    acao: str,
    recurso: str,
    detalhes: dict,
) -> None:
    session.add(
        EventoAuditoria(
            organizacao_id=usuario.organizacao_id,
            actor_id=usuario.id,
            ator=usuario.ator,
            acao=acao[:20],
            recurso=recurso[:180],
            sucesso=True,
            status_http=200,
            ip_hash=hash_ip(cliente_ip(request)),
            detalhes=detalhes,
        )
    )


async def _vincular_ids(
    session: AsyncSession,
    request: Request,
    usuario: UsuarioAutenticado,
    processo_ids: list[int],
    dados: VinculoBase,
    *,
    origem: str,
    procurador_origem: str | None = None,
) -> dict:
    ids = list(dict.fromkeys(processo_ids))
    processos_existentes = set(
        (
            await session.execute(
                select(Processo.id).where(
                    Processo.id.in_(ids), Processo.tipo == TipoProcesso.MARCA
                )
            )
        ).scalars().all()
    )
    ja_vinculados = set(
        (
            await session.execute(
                select(ProcessoMonitorado.processo_id).where(
                    ProcessoMonitorado.organizacao_id == usuario.organizacao_id,
                    ProcessoMonitorado.processo_id.in_(processos_existentes),
                )
            )
        ).scalars().all()
    )
    empresa = await _empresa(
        session, usuario, dados.empresa_id, dados.empresa_nome
    )
    await _validar_responsavel(session, usuario, dados.responsavel_id)
    novos = processos_existentes - ja_vinculados
    for processo_id in novos:
        session.add(
            ProcessoMonitorado(
                organizacao_id=usuario.organizacao_id,
                processo_id=processo_id,
                empresa_id=empresa.id if empresa else None,
                responsavel_id=dados.responsavel_id,
                status="ativo",
                origem=origem,
                procurador_origem=procurador_origem,
                observacoes=dados.observacoes,
                vinculado_por=usuario.ator,
            )
        )
    resultado = {
        "encontrados": len(ids),
        "vinculados": len(novos),
        "ja_vinculados": len(ja_vinculados),
        "nao_encontrados": len(set(ids) - processos_existentes),
        "empresa": empresa.nome if empresa else None,
    }
    _auditar(
        session,
        request,
        usuario,
        "vincular_processos",
        f"carteira:{origem}",
        {**resultado, "procurador": procurador_origem},
    )
    await session.commit()
    return resultado


@router.get("")
async def listar_carteira(
    session: SessionDep,
    usuario: ViewDep,
    busca: Annotated[str | None, Query(max_length=150)] = None,
    status: Annotated[str | None, Query(max_length=20)] = None,
    etapa: Annotated[EtapaKanban | None, Query()] = None,
    limite: Annotated[int, Query(ge=1, le=20)] = 20,
    deslocamento: Annotated[int, Query(ge=0)] = 0,
    situacao_inpi: Annotated[str | None, Query(max_length=30)] = None,
) -> dict:
    situacao_inpi = _validar_grupo_situacao_inpi(situacao_inpi)
    resumo_linhas = (
        await session.execute(
            select(ProcessoMonitorado.status, func.count())
            .where(ProcessoMonitorado.organizacao_id == usuario.organizacao_id)
            .group_by(ProcessoMonitorado.status)
        )
    ).all()
    resumo = {chave: int(total_status) for chave, total_status in resumo_linhas}
    grupo_resumo = _expressao_grupo_situacao_inpi()
    resumo_situacoes_linhas = (
        await session.execute(
            select(grupo_resumo.label("grupo"), func.count())
            .select_from(ProcessoMonitorado)
            .join(Processo, Processo.id == ProcessoMonitorado.processo_id)
            .where(ProcessoMonitorado.organizacao_id == usuario.organizacao_id)
            .group_by(grupo_resumo)
        )
    ).all()
    resumo_situacoes = {
        chave: int(total_situacao) for chave, total_situacao in resumo_situacoes_linhas
    }
    filtros = [ProcessoMonitorado.organizacao_id == usuario.organizacao_id]
    if status:
        filtros.append(ProcessoMonitorado.status == status)
    if etapa:
        filtros.append(ProcessoMonitorado.etapa_kanban == etapa)
    if situacao_inpi:
        filtros.append(_expressao_grupo_situacao_inpi() == situacao_inpi)
    if busca:
        termo = f"%{_normalizar_busca(busca)}%"
        filtros.append(
            or_(
                Processo.numero_normalizado.ilike(f"%{normalizar_numero_processo(busca)}%"),
                func.immutable_unaccent(func.lower(Processo.titulo)).ilike(termo),
                _expressao_procurador().ilike(termo),
                func.immutable_unaccent(func.lower(EmpresaCRM.nome)).ilike(termo),
            )
        )
    base = (
        select(ProcessoMonitorado)
        .join(Processo, Processo.id == ProcessoMonitorado.processo_id)
        .outerjoin(EmpresaCRM, EmpresaCRM.id == ProcessoMonitorado.empresa_id)
        .where(*filtros)
    )
    prioridade_situacao = case(
        (_expressao_grupo_situacao_inpi() == "deferido", 1),
        (_expressao_grupo_situacao_inpi() == "em_tramitacao", 2),
        (_expressao_grupo_situacao_inpi() == "registrado", 3),
        (_expressao_grupo_situacao_inpi() == "indeferido", 4),
        (_expressao_grupo_situacao_inpi() == "encerrado", 5),
        else_=6,
    )
    total = int(
        (await session.execute(select(func.count()).select_from(base.subquery()))).scalar_one()
    )
    ultima_rpi = (
        select(Movimentacao.numero_rpi)
        .where(Movimentacao.processo_id == Processo.id)
        .order_by(Movimentacao.data_rpi.desc(), Movimentacao.id.desc())
        .limit(1)
        .scalar_subquery()
    )
    ultima_data = (
        select(Movimentacao.data_rpi)
        .where(Movimentacao.processo_id == Processo.id)
        .order_by(Movimentacao.data_rpi.desc(), Movimentacao.id.desc())
        .limit(1)
        .scalar_subquery()
    )
    ultima_descricao = (
        select(Movimentacao.descricao)
        .where(Movimentacao.processo_id == Processo.id)
        .order_by(Movimentacao.data_rpi.desc(), Movimentacao.id.desc())
        .limit(1)
        .scalar_subquery()
    )
    linhas = (
        await session.execute(
            select(
                ProcessoMonitorado,
                Processo,
                EmpresaCRM.nome,
                UsuarioOperacoes.nome,
                ultima_rpi,
                ultima_data,
                ultima_descricao,
            )
            .join(Processo, Processo.id == ProcessoMonitorado.processo_id)
            .outerjoin(EmpresaCRM, EmpresaCRM.id == ProcessoMonitorado.empresa_id)
            .outerjoin(
                UsuarioOperacoes, UsuarioOperacoes.id == ProcessoMonitorado.responsavel_id
            )
            .where(*filtros)
            .order_by(prioridade_situacao, ProcessoMonitorado.atualizado_em.desc(), Processo.numero)
            .limit(limite)
            .offset(deslocamento)
        )
    ).all()
    return {
        "total": total,
        "limite": limite,
        "deslocamento": deslocamento,
        "resumo": {
            "total": sum(resumo.values()),
            "pausados": resumo.get("pausado", 0),
            "deferidos": resumo_situacoes.get("deferido", 0),
            "arquivados_extintos": resumo_situacoes.get("encerrado", 0),
            "indeferidos": resumo_situacoes.get("indeferido", 0),
            "registros_concedidos": resumo_situacoes.get("registrado", 0),
            "em_tramitacao": resumo_situacoes.get("em_tramitacao", 0),
        },
        "itens": [
            {
                "id": monitorado.id,
                "numero": processo.numero,
                "titulo": processo.titulo,
                "titulo_exibicao": _titulo_exibicao(processo),
                "situacao": processo.situacao,
                "situacao_normalizada": processo.situacao_normalizada,
                "grupo_situacao_inpi": _grupo_situacao_valor(
                    processo.situacao_normalizada
                ),
                "grupo_situacao_inpi_nome": _nome_grupo_situacao_inpi(
                    _grupo_situacao_valor(processo.situacao_normalizada)
                ),
                "data_deposito": processo.data_deposito,
                "procurador": processo.procurador,
                "fonte": processo.fonte,
                "processo_atualizado_em": processo.atualizado_em,
                "status": monitorado.status,
                "etapa_kanban": monitorado.etapa_kanban,
                "etapa_atualizada_em": monitorado.etapa_atualizada_em,
                "etapa_atualizada_por": monitorado.etapa_atualizada_por,
                "origem": monitorado.origem,
                "empresa_id": monitorado.empresa_id,
                "empresa": empresa_nome,
                "responsavel_id": monitorado.responsavel_id,
                "responsavel": responsavel_nome,
                "observacoes": monitorado.observacoes,
                "criado_em": monitorado.criado_em,
                "ultima_movimentacao": (
                    {
                        "numero_rpi": numero_rpi,
                        "data": data_rpi,
                        "descricao": descricao,
                    }
                    if numero_rpi is not None
                    else None
                ),
            }
            for (
                monitorado,
                processo,
                empresa_nome,
                responsavel_nome,
                numero_rpi,
                data_rpi,
                descricao,
            ) in linhas
        ],
        "tem_mais": deslocamento + len(linhas) < total,
    }


@router.get("/kanban")
async def listar_kanban(
    session: SessionDep,
    usuario: ViewDep,
    busca: Annotated[str | None, Query(max_length=150)] = None,
    status: Annotated[str | None, Query(max_length=20)] = None,
    situacao_inpi: Annotated[str | None, Query(max_length=30)] = None,
) -> dict:
    """Entrega no máximo 20 cartões por etapa e o total real de cada coluna."""
    situacao_inpi = _validar_grupo_situacao_inpi(situacao_inpi)
    filtros = [ProcessoMonitorado.organizacao_id == usuario.organizacao_id]
    if status:
        filtros.append(ProcessoMonitorado.status == status)
    if situacao_inpi:
        filtros.append(_expressao_grupo_situacao_inpi() == situacao_inpi)
    if busca:
        termo = f"%{_normalizar_busca(busca)}%"
        filtros.append(
            or_(
                Processo.numero_normalizado.ilike(
                    f"%{normalizar_numero_processo(busca)}%"
                ),
                func.immutable_unaccent(func.lower(Processo.titulo)).ilike(termo),
                _expressao_procurador().ilike(termo),
                func.immutable_unaccent(func.lower(EmpresaCRM.nome)).ilike(termo),
            )
        )

    contagens = dict(
        (
            await session.execute(
                select(ProcessoMonitorado.etapa_kanban, func.count())
                .join(Processo, Processo.id == ProcessoMonitorado.processo_id)
                .outerjoin(EmpresaCRM, EmpresaCRM.id == ProcessoMonitorado.empresa_id)
                .where(*filtros)
                .group_by(ProcessoMonitorado.etapa_kanban)
            )
        ).all()
    )


    ultima_rpi = (
        select(Movimentacao.numero_rpi)
        .where(Movimentacao.processo_id == Processo.id)
        .order_by(Movimentacao.data_rpi.desc(), Movimentacao.id.desc())
        .limit(1)
        .scalar_subquery()
    )
    ultima_data = (
        select(Movimentacao.data_rpi)
        .where(Movimentacao.processo_id == Processo.id)
        .order_by(Movimentacao.data_rpi.desc(), Movimentacao.id.desc())
        .limit(1)
        .scalar_subquery()
    )
    ultima_descricao = (
        select(Movimentacao.descricao)
        .where(Movimentacao.processo_id == Processo.id)
        .order_by(Movimentacao.data_rpi.desc(), Movimentacao.id.desc())
        .limit(1)
        .scalar_subquery()
    )
    colunas = []
    for chave, titulo in ETAPAS_KANBAN:
        linhas = (
            await session.execute(
                select(
                    ProcessoMonitorado,
                    Processo,
                    EmpresaCRM.nome,
                    UsuarioOperacoes.nome,
                    ultima_rpi,
                    ultima_data,
                    ultima_descricao,
                )
                .join(Processo, Processo.id == ProcessoMonitorado.processo_id)
                .outerjoin(EmpresaCRM, EmpresaCRM.id == ProcessoMonitorado.empresa_id)
                .outerjoin(
                    UsuarioOperacoes,
                    UsuarioOperacoes.id == ProcessoMonitorado.responsavel_id,
                )
                .where(*filtros, ProcessoMonitorado.etapa_kanban == chave)
                .order_by(
                    ProcessoMonitorado.ordem_kanban,
                    ProcessoMonitorado.etapa_atualizada_em.desc(),
                    ProcessoMonitorado.id.desc(),
                )
                .limit(20)
            )
        ).all()
        itens = []
        for (
            monitorado,
            processo,
            empresa_nome,
            responsavel_nome,
            numero_rpi,
            data_rpi,
            descricao,
        ) in linhas:
            itens.append(
                {
                    "id": monitorado.id,
                    "numero": processo.numero,
                    "titulo": processo.titulo,
                    "titulo_exibicao": _titulo_exibicao(processo),
                    "situacao": processo.situacao,
                    "data_deposito": processo.data_deposito,
                    "status": monitorado.status,
                    "etapa_kanban": monitorado.etapa_kanban,
                    "empresa": empresa_nome,
                    "responsavel": responsavel_nome,
                    "ultima_movimentacao": (
                        {
                            "numero_rpi": numero_rpi,
                            "data": data_rpi,
                            "descricao": descricao,
                        }
                        if numero_rpi is not None
                        else None
                    ),
                }
            )
        total = int(contagens.get(chave, 0))
        colunas.append(
            {
                "chave": chave,
                "titulo": titulo,
                "total": total,
                "limite": 20,
                "tem_mais": total > len(itens),
                "itens": itens,
            }
        )
    return {"total": sum(int(valor) for valor in contagens.values()), "colunas": colunas}


@router.get("/kanban-inpi")
async def listar_kanban_inpi(
    session: SessionDep,
    usuario: ViewDep,
    busca: Annotated[str | None, Query(max_length=150)] = None,
    status: Annotated[str | None, Query(max_length=20)] = None,
    situacao_inpi: Annotated[str | None, Query(max_length=30)] = None,
) -> dict:
    """Organiza automaticamente a carteira pela situação oficial publicada na RPI."""
    situacao_inpi = _validar_grupo_situacao_inpi(situacao_inpi)
    grupo = _expressao_grupo_situacao_inpi()
    filtros = [ProcessoMonitorado.organizacao_id == usuario.organizacao_id]
    if status:
        filtros.append(ProcessoMonitorado.status == status)
    if situacao_inpi:
        filtros.append(grupo == situacao_inpi)
    if busca:
        termo = f"%{_normalizar_busca(busca)}%"
        filtros.append(
            or_(
                Processo.numero_normalizado.ilike(
                    f"%{normalizar_numero_processo(busca)}%"
                ),
                func.immutable_unaccent(func.lower(Processo.titulo)).ilike(termo),
                _expressao_procurador().ilike(termo),
                func.immutable_unaccent(func.lower(EmpresaCRM.nome)).ilike(termo),
            )
        )

    contagens = dict(
        (
            await session.execute(
                select(grupo.label("grupo"), func.count())
                .join(Processo, Processo.id == ProcessoMonitorado.processo_id)
                .outerjoin(EmpresaCRM, EmpresaCRM.id == ProcessoMonitorado.empresa_id)
                .where(*filtros)
                .group_by(grupo)
            )
        ).all()
    )
    ultima_rpi = (
        select(Movimentacao.numero_rpi)
        .where(Movimentacao.processo_id == Processo.id)
        .order_by(Movimentacao.data_rpi.desc(), Movimentacao.id.desc())
        .limit(1)
        .scalar_subquery()
    )
    ultima_data = (
        select(Movimentacao.data_rpi)
        .where(Movimentacao.processo_id == Processo.id)
        .order_by(Movimentacao.data_rpi.desc(), Movimentacao.id.desc())
        .limit(1)
        .scalar_subquery()
    )
    ultima_descricao = (
        select(Movimentacao.descricao)
        .where(Movimentacao.processo_id == Processo.id)
        .order_by(Movimentacao.data_rpi.desc(), Movimentacao.id.desc())
        .limit(1)
        .scalar_subquery()
    )
    colunas = []
    for chave, titulo in GRUPOS_SITUACAO_INPI:
        linhas = (
            await session.execute(
                select(
                    ProcessoMonitorado,
                    Processo,
                    EmpresaCRM.nome,
                    UsuarioOperacoes.nome,
                    ultima_rpi,
                    ultima_data,
                    ultima_descricao,
                )
                .join(Processo, Processo.id == ProcessoMonitorado.processo_id)
                .outerjoin(EmpresaCRM, EmpresaCRM.id == ProcessoMonitorado.empresa_id)
                .outerjoin(
                    UsuarioOperacoes,
                    UsuarioOperacoes.id == ProcessoMonitorado.responsavel_id,
                )
                .where(*filtros, grupo == chave)
                .order_by(
                    ProcessoMonitorado.atualizado_em.desc(),
                    ProcessoMonitorado.id.desc(),
                )
                .limit(20)
            )
        ).all()
        itens = [
            {
                "id": monitorado.id,
                "numero": processo.numero,
                "titulo": processo.titulo,
                "titulo_exibicao": _titulo_exibicao(processo),
                "situacao": processo.situacao,
                "situacao_normalizada": processo.situacao_normalizada,
                "grupo_situacao_inpi": chave,
                "data_deposito": processo.data_deposito,
                "status": monitorado.status,
                "etapa_kanban": monitorado.etapa_kanban,
                "empresa": empresa_nome,
                "responsavel": responsavel_nome,
                "ultima_movimentacao": (
                    {
                        "numero_rpi": numero_rpi,
                        "data": data_rpi,
                        "descricao": descricao,
                    }
                    if numero_rpi is not None
                    else None
                ),
            }
            for (
                monitorado,
                processo,
                empresa_nome,
                responsavel_nome,
                numero_rpi,
                data_rpi,
                descricao,
            ) in linhas
        ]
        total = int(contagens.get(chave, 0))
        colunas.append(
            {
                "chave": chave,
                "titulo": titulo,
                "total": total,
                "limite": 20,
                "tem_mais": total > len(itens),
                "itens": itens,
            }
        )
    return {
        "total": sum(int(valor) for valor in contagens.values()),
        "modo": "inpi",
        "fonte": "Situação consolidada a partir da última publicação na RPI",
        "colunas": colunas,
    }


@router.get("/empresas")
async def listar_empresas(
    session: SessionDep,
    usuario: ViewDep,
    busca: Annotated[str | None, Query(max_length=100)] = None,
) -> list[dict]:
    consulta = select(EmpresaCRM).where(
        EmpresaCRM.organizacao_id == usuario.organizacao_id
    )
    if busca:
        consulta = consulta.where(
            func.immutable_unaccent(func.lower(EmpresaCRM.nome)).ilike(
                f"%{_normalizar_busca(busca)}%"
            )
        )
    empresas = (
        await session.execute(consulta.order_by(EmpresaCRM.nome).limit(100))
    ).scalars().all()
    return [{"id": empresa.id, "nome": empresa.nome} for empresa in empresas]


@router.get("/responsaveis")
async def listar_responsaveis(
    session: SessionDep, usuario: ViewDep
) -> list[dict]:
    pessoas = (
        await session.execute(
            select(UsuarioOperacoes)
            .where(
                UsuarioOperacoes.organizacao_id == usuario.organizacao_id,
                UsuarioOperacoes.ativo.is_(True),
            )
            .order_by(UsuarioOperacoes.nome)
        )
    ).scalars().all()
    return [{"id": pessoa.id, "nome": pessoa.nome} for pessoa in pessoas]


@router.get("/procuradores")
async def sugerir_procuradores(
    session: SessionDep,
    _usuario: ViewDep,
    busca: Annotated[str, Query(min_length=2, max_length=120)],
) -> list[str]:
    termo = f"%{_normalizar_busca(busca)}%"
    nomes = (
        await session.execute(
            select(Processo.procurador)
            .where(
                Processo.tipo == TipoProcesso.MARCA,
                Processo.procurador.is_not(None),
                _expressao_procurador().ilike(termo),
            )
            .distinct()
            .order_by(Processo.procurador)
            .limit(20)
        )
    ).scalars().all()
    return list(nomes)


def _filtro_procurador(procurador: str, modo: str):
    normalizado = _normalizar_busca(procurador)
    expressao = _expressao_procurador()
    if modo == "exato":
        return expressao == normalizado
    if modo == "variacoes":
        return and_(*(expressao.ilike(f"%{termo}%") for termo in normalizado.split()))
    return expressao.ilike(f"%{normalizado}%")


@router.get("/buscar-procurador")
async def buscar_por_procurador(
    session: SessionDep,
    usuario: ViewDep,
    procurador: Annotated[str, Query(min_length=2, max_length=300)],
    modo: Literal["exato", "contem", "variacoes"] = "exato",
    limite: Annotated[int, Query(ge=1, le=100)] = 50,
    deslocamento: Annotated[int, Query(ge=0)] = 0,
) -> dict:
    filtro = _filtro_procurador(procurador, modo)
    filtros = [Processo.tipo == TipoProcesso.MARCA, filtro]
    total, titulares = (
        await session.execute(
            select(
                func.count(func.distinct(Processo.id)),
                func.count(func.distinct(processo_titulares.c.titular_id)),
            )
            .select_from(Processo)
            .outerjoin(
                processo_titulares,
                processo_titulares.c.processo_id == Processo.id,
            )
            .where(*filtros)
        )
    ).one()
    primeira_rpi, ultima_rpi, rpis_importadas = (
        await session.execute(
            select(
                func.min(RpiImportacao.numero_rpi),
                func.max(RpiImportacao.numero_rpi),
                func.count(),
            ).where(RpiImportacao.tipo == "marca")
        )
    ).one()
    vinculo = (
        select(ProcessoMonitorado.id)
        .where(
            ProcessoMonitorado.organizacao_id == usuario.organizacao_id,
            ProcessoMonitorado.processo_id == Processo.id,
        )
        .limit(1)
        .scalar_subquery()
    )
    linhas = (
        await session.execute(
            select(Processo, vinculo)
            .where(*filtros)
            .order_by(Processo.atualizado_em.desc(), Processo.numero)
            .limit(limite)
            .offset(deslocamento)
        )
    ).all()
    variacoes: dict[str, int] = {}
    for processo, _ in linhas:
        if processo.procurador:
            variacoes[processo.procurador] = variacoes.get(processo.procurador, 0) + 1
    return {
        "total": int(total or 0),
        "total_processos": int(total or 0),
        "total_titulares": int(titulares or 0),
        "limite": limite,
        "deslocamento": deslocamento,
        "modo": modo,
        "variacoes": [
            {"nome": nome, "processos_na_pagina": quantidade}
            for nome, quantidade in sorted(
                variacoes.items(),
                key=lambda item: (-item[1], item[0].casefold()),
            )
        ],
        "cobertura": {
            "primeira_rpi": primeira_rpi,
            "ultima_rpi": ultima_rpi,
            "rpis_importadas": int(rpis_importadas or 0),
            "aviso": (
                "O resultado depende do nome do procurador publicado nas RPIs e pode não "
                "representar toda a carteira histórica existente no INPI."
            ),
        },
        "itens": [
            {
                "processo_id": processo.id,
                "numero": processo.numero,
                "titulo": processo.titulo,
                "titulo_exibicao": _titulo_exibicao(processo),
                "data_deposito": processo.data_deposito,
                "situacao": processo.situacao,
                "procurador": processo.procurador,
                "fonte": processo.fonte,
                "monitorado_id": monitorado_id,
            }
            for processo, monitorado_id in linhas
        ],
    }


@router.post("/manual", status_code=201)
async def cadastrar_manual(
    dados: CadastroManual,
    request: Request,
    session: SessionDep,
    usuario: ManageDep,
) -> dict:
    processo = (
        await session.execute(
            select(Processo).where(
                Processo.numero_normalizado == normalizar_numero_processo(dados.numero),
                Processo.tipo == TipoProcesso.MARCA,
            )
        )
    ).scalar_one_or_none()
    if processo is None:
        raise HTTPException(
            404,
            "Processo não localizado na base RPI. Sincronize as revistas antes de cadastrar.",
        )
    resultado = await _vincular_ids(
        session, request, usuario, [processo.id], dados, origem="manual"
    )
    return {**resultado, "numero": processo.numero}


COLUNAS_NUMERO = {"numero", "processo", "numeroprocesso", "nprocesso", "registro"}
COLUNAS_EMPRESA = {"empresa", "cliente", "titular", "razaosocial"}
COLUNAS_PROCURADOR = {"procurador", "agente", "escritorio"}
COLUNAS_OBS = {"observacoes", "observacao", "obs", "notas"}
TAMANHO_MAXIMO_IMPORTACAO = 5_000_000


def _chave_coluna(texto: str) -> str:
    sem_acentos = "".join(
        c for c in unicodedata.normalize("NFKD", texto or "") if not unicodedata.combining(c)
    )
    return "".join(ch for ch in sem_acentos.lower() if ch.isalnum())


def _ler_planilha(conteudo: bytes, filename: str) -> list[dict[str, str]]:
    """Lê CSV ou XLSX e devolve uma lista de registros com chaves normalizadas."""
    nome = (filename or "").lower()
    linhas: list[list[str]] = []
    if nome.endswith((".xlsx", ".xlsm")):
        try:
            wb = load_workbook(io.BytesIO(conteudo), read_only=True, data_only=True)
        except Exception as exc:  # noqa: BLE001 - arquivo inválido enviado pelo usuário
            raise HTTPException(400, "Planilha Excel inválida ou corrompida.") from exc
        planilha = wb.active
        for linha in planilha.iter_rows(values_only=True):
            linhas.append(["" if celula is None else str(celula).strip() for celula in linha])
        wb.close()
    else:
        texto = None
        for codificacao in ("utf-8-sig", "latin-1"):
            try:
                texto = conteudo.decode(codificacao)
                break
            except UnicodeDecodeError:
                continue
        if texto is None:
            raise HTTPException(400, "Não foi possível ler o arquivo (codificação não suportada).")
        delimitador = ";" if texto.count(";") > texto.count(",") else ","
        for linha in csv.reader(io.StringIO(texto), delimiter=delimitador):
            linhas.append([campo.strip() for campo in linha])

    linhas = [linha for linha in linhas if any(linha)]
    if len(linhas) < 2:
        return []
    cabecalho = [_chave_coluna(coluna) for coluna in linhas[0]]
    return [
        {cabecalho[i]: (linha[i] if i < len(linha) else "") for i in range(len(cabecalho))}
        for linha in linhas[1:]
    ]


def _valor(registro: dict[str, str], chaves: set[str]) -> str | None:
    for chave, valor in registro.items():
        if chave in chaves and valor.strip():
            return valor.strip()
    return None


@router.post("/importar", status_code=201)
async def importar_carteira(
    request: Request,
    session: SessionDep,
    usuario: ManageDep,
    arquivo: Annotated[UploadFile, File()],
    responsavel_id: Annotated[int | None, Form()] = None,
) -> dict:
    """Importa uma carteira (CSV/XLSX) para os processos monitorados.

    Colunas reconhecidas (cabeçalho, sem acento/maiúsculas): numero (obrigatória),
    empresa, procurador, observacoes. Números não localizados na base RPI são
    reportados; duplicados já monitorados são ignorados.
    """
    conteudo = await arquivo.read()
    if not conteudo:
        raise HTTPException(400, "Arquivo vazio.")
    if len(conteudo) > TAMANHO_MAXIMO_IMPORTACAO:
        raise HTTPException(413, "Arquivo muito grande (máximo 5 MB).")
    registros = _ler_planilha(conteudo, arquivo.filename or "")
    if not registros:
        raise HTTPException(
            400,
            "Planilha vazia ou sem cabeçalho reconhecível. Inclua uma coluna 'numero'.",
        )
    await _validar_responsavel(session, usuario, responsavel_id)

    linhas_validas: list[tuple[str, str, dict[str, str]]] = []
    sem_numero = 0
    for registro in registros:
        numero = _valor(registro, COLUNAS_NUMERO)
        if not numero:
            sem_numero += 1
            continue
        linhas_validas.append((numero, normalizar_numero_processo(numero), registro))
    if not linhas_validas:
        raise HTTPException(400, "Nenhuma linha com número de processo foi encontrada.")

    normalizados = {norm for _, norm, _ in linhas_validas}
    processos = {
        row.numero_normalizado: row
        for row in (
            await session.execute(
                select(Processo).where(
                    Processo.numero_normalizado.in_(normalizados),
                    Processo.tipo == TipoProcesso.MARCA,
                )
            )
        ).scalars()
    }
    ja_monitorados = set(
        (
            await session.execute(
                select(ProcessoMonitorado.processo_id).where(
                    ProcessoMonitorado.organizacao_id == usuario.organizacao_id,
                    ProcessoMonitorado.processo_id.in_(
                        processo.id for processo in processos.values()
                    ),
                )
            )
        ).scalars().all()
    )

    empresas_cache: dict[str, int | None] = {}
    processados: set[int] = set()
    vinculados = 0
    ja_vinculados = 0
    nao_encontrados: list[str] = []
    for numero, norm, registro in linhas_validas:
        processo = processos.get(norm)
        if processo is None:
            nao_encontrados.append(numero)
            continue
        if processo.id in ja_monitorados or processo.id in processados:
            ja_vinculados += 1
            continue
        empresa_nome = _valor(registro, COLUNAS_EMPRESA)
        if empresa_nome and empresa_nome not in empresas_cache:
            empresa = await obter_ou_criar_empresa(session, usuario.organizacao_id, empresa_nome)
            empresas_cache[empresa_nome] = empresa.id if empresa else None
        session.add(
            ProcessoMonitorado(
                organizacao_id=usuario.organizacao_id,
                processo_id=processo.id,
                empresa_id=empresas_cache.get(empresa_nome) if empresa_nome else None,
                responsavel_id=responsavel_id,
                status="ativo",
                origem="importacao",
                procurador_origem=_valor(registro, COLUNAS_PROCURADOR),
                observacoes=_valor(registro, COLUNAS_OBS),
                vinculado_por=usuario.ator,
            )
        )
        processados.add(processo.id)
        vinculados += 1

    resultado = {
        "total_linhas": len(registros),
        "vinculados": vinculados,
        "ja_vinculados": ja_vinculados,
        "nao_encontrados": len(nao_encontrados),
        "sem_numero": sem_numero,
        "empresas_associadas": len([v for v in empresas_cache.values() if v]),
        "exemplos_nao_encontrados": nao_encontrados[:20],
    }
    _auditar(
        session,
        request,
        usuario,
        "importar_carteira",
        f"carteira:importacao:{arquivo.filename}",
        {k: v for k, v in resultado.items() if k != "exemplos_nao_encontrados"},
    )
    await session.commit()
    return resultado


@router.post("/vincular-lote", status_code=201)
async def vincular_lote(
    dados: VinculoLote,
    request: Request,
    session: SessionDep,
    usuario: ManageDep,
) -> dict:
    return await _vincular_ids(
        session, request, usuario, dados.processo_ids, dados, origem="selecao_procurador"
    )


@router.post("/vincular-procurador", status_code=201)
async def vincular_todos_do_procurador(
    dados: VinculoProcurador,
    request: Request,
    session: SessionDep,
    usuario: ManageDep,
) -> dict:
    ids = list(
        (
            await session.execute(
                select(Processo.id)
                .where(
                    Processo.tipo == TipoProcesso.MARCA,
                    _filtro_procurador(dados.procurador, dados.modo),
                )
                .order_by(Processo.id)
                .limit(dados.maximo)
            )
        ).scalars().all()
    )
    if not ids:
        raise HTTPException(404, "Nenhum processo encontrado para o procurador informado")
    return await _vincular_ids(
        session,
        request,
        usuario,
        ids,
        dados,
        origem="procurador",
        procurador_origem=dados.procurador,
    )


@router.get("/{monitorado_id}/historico-kanban")
async def historico_kanban(
    monitorado_id: int,
    session: SessionDep,
    usuario: ViewDep,
) -> list[dict]:
    existe = await session.scalar(
        select(ProcessoMonitorado.id).where(
            ProcessoMonitorado.id == monitorado_id,
            ProcessoMonitorado.organizacao_id == usuario.organizacao_id,
        )
    )
    if existe is None:
        raise HTTPException(404, "Processo monitorado não encontrado")
    eventos = (
        await session.execute(
            select(HistoricoEtapaCarteira)
            .where(
                HistoricoEtapaCarteira.organizacao_id == usuario.organizacao_id,
                HistoricoEtapaCarteira.processo_monitorado_id == monitorado_id,
            )
            .order_by(HistoricoEtapaCarteira.criado_em.desc())
            .limit(50)
        )
    ).scalars().all()
    return [
        {
            "id": evento.id,
            "etapa_anterior": evento.etapa_anterior,
            "etapa_nova": evento.etapa_nova,
            "movido_por": evento.movido_por,
            "criado_em": evento.criado_em,
        }
        for evento in eventos
    ]


@router.patch("/{monitorado_id}")
async def atualizar_monitoramento(
    monitorado_id: int,
    dados: AtualizacaoMonitoramento,
    request: Request,
    session: SessionDep,
    usuario: ManageDep,
) -> dict:
    monitorado = (
        await session.execute(
            select(ProcessoMonitorado).where(
                ProcessoMonitorado.id == monitorado_id,
                ProcessoMonitorado.organizacao_id == usuario.organizacao_id,
            )
        )
    ).scalar_one_or_none()
    if monitorado is None:
        raise HTTPException(404, "Processo monitorado não encontrado")
    antes = {
        "status": monitorado.status,
        "empresa_id": monitorado.empresa_id,
        "responsavel_id": monitorado.responsavel_id,
        "etapa_kanban": monitorado.etapa_kanban,
    }
    if dados.status is not None:
        monitorado.status = dados.status
    if dados.empresa_id is not None or dados.empresa_nome:
        empresa = await _empresa(
            session, usuario, dados.empresa_id, dados.empresa_nome
        )
        monitorado.empresa_id = empresa.id if empresa else None
    if dados.remover_responsavel:
        monitorado.responsavel_id = None
    elif dados.responsavel_id is not None:
        await _validar_responsavel(session, usuario, dados.responsavel_id)
        monitorado.responsavel_id = dados.responsavel_id
    if dados.observacoes is not None:
        monitorado.observacoes = dados.observacoes.strip() or None
    if dados.procurador is not None:
        # procurador é dado compartilhado do Processo (não do vínculo). Permitimos
        # corrigi-lo a partir da carteira porque muitas marcas do BADEPI vêm sem ele.
        processo = await session.get(Processo, monitorado.processo_id)
        if processo is not None:
            processo.procurador = dados.procurador or None
    if dados.etapa_kanban is not None and dados.etapa_kanban != monitorado.etapa_kanban:
        etapa_anterior = monitorado.etapa_kanban
        monitorado.etapa_kanban = dados.etapa_kanban
        monitorado.ordem_kanban = 0
        monitorado.etapa_atualizada_em = datetime.now(UTC)
        monitorado.etapa_atualizada_por = usuario.ator
        session.add(
            HistoricoEtapaCarteira(
                organizacao_id=usuario.organizacao_id,
                processo_monitorado_id=monitorado.id,
                etapa_anterior=etapa_anterior,
                etapa_nova=dados.etapa_kanban,
                movido_por=usuario.ator,
            )
        )
    monitorado.atualizado_em = datetime.now(UTC)
    _auditar(
        session,
        request,
        usuario,
        "atualizar_carteira",
        f"processo-monitorado:{monitorado.id}",
        {
            "antes": antes,
            "depois": {
                "status": monitorado.status,
                "empresa_id": monitorado.empresa_id,
                "responsavel_id": monitorado.responsavel_id,
                "etapa_kanban": monitorado.etapa_kanban,
            },
        },
    )
    await session.commit()
    return {"status": "ok", "id": monitorado.id}


@router.post("/{monitorado_id}/atualizar")
async def atualizar_status_processo(
    monitorado_id: int,
    request: Request,
    session: SessionDep,
    usuario: ManageDep,
) -> dict:
    """Reconsolida a situação a partir dos despachos já sincronizados (feed RPI) e,
    se o processo ainda não tiver cliente, cadastra a empresa a partir do titular."""
    monitorado = (
        await session.execute(
            select(ProcessoMonitorado).where(
                ProcessoMonitorado.id == monitorado_id,
                ProcessoMonitorado.organizacao_id == usuario.organizacao_id,
            )
        )
    ).scalar_one_or_none()
    if monitorado is None:
        raise HTTPException(404, "Processo monitorado não encontrado")

    await consolidar_situacao(session, monitorado.processo_id)
    processo = await session.get(Processo, monitorado.processo_id)

    cliente_cadastrado = None
    if monitorado.empresa_id is None:
        titular_nome = (
            await session.execute(
                select(Titular.nome)
                .join(processo_titulares, processo_titulares.c.titular_id == Titular.id)
                .where(processo_titulares.c.processo_id == monitorado.processo_id)
                .order_by(Titular.nome)
                .limit(1)
            )
        ).scalar_one_or_none()
        empresa = await obter_ou_criar_empresa(session, usuario.organizacao_id, titular_nome)
        if empresa is not None:
            monitorado.empresa_id = empresa.id
            cliente_cadastrado = empresa.nome

    monitorado.atualizado_em = datetime.now(UTC)
    _auditar(
        session,
        request,
        usuario,
        "atualizar_status_carteira",
        f"processo-monitorado:{monitorado.id}",
        {
            "situacao": processo.situacao if processo else None,
            "cliente_cadastrado": cliente_cadastrado,
        },
    )
    await session.commit()
    return {
        "status": "ok",
        "situacao": processo.situacao if processo else None,
        "situacao_normalizada": processo.situacao_normalizada if processo else None,
        "relevancia_situacao": processo.relevancia_situacao if processo else None,
        "cliente_cadastrado": cliente_cadastrado,
    }
